"""
Advanced Document Preprocessing Service
다양한 문서 형식에 대한 고급 전처리 기능 제공
- 텍스트 추출 실패 시 여러 fallback 방법 시도
- 문서 구조 분석 및 메타데이터 추출
- 다양한 인코딩 및 포맷 지원
"""

import os
import io
import json
import uuid
import tempfile
import logging
from typing import List, Dict, Any, Optional, Tuple, Union
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

logger = logging.getLogger(__name__)

@dataclass
class PreprocessingResult:
    """전처리 결과"""
    success: bool
    extracted_text: str
    method_used: str
    metadata: Dict[str, Any]
    processing_time: float
    char_count: int
    error_message: Optional[str] = None
    fallback_attempts: List[str] = None

class AdvancedDocumentPreprocessor:
    """고급 문서 전처리기"""

    def __init__(self):
        self.supported_formats = ['.docx', '.doc', '.pdf', '.pptx', '.ppt', '.xlsx', '.xls', '.txt', '.md']

    async def preprocess_document(self, file_path: str, filename: str) -> PreprocessingResult:
        """문서 전처리 - 다양한 방법으로 텍스트 추출 시도"""
        start_time = datetime.now()

        # 파일 확장자 확인
        file_ext = Path(filename).suffix.lower()

        logger.info(f"🔧 [PREPROCESS] Starting advanced preprocessing for {filename} ({file_ext})")

        try:
            if file_ext in ['.docx', '.doc']:
                result = await self._preprocess_word_document(file_path, filename)
            elif file_ext == '.pdf':
                result = await self._preprocess_pdf_document(file_path, filename)
            elif file_ext in ['.pptx', '.ppt']:
                result = await self._preprocess_powerpoint_document(file_path, filename)
            elif file_ext in ['.xlsx', '.xls']:
                result = await self._preprocess_excel_document(file_path, filename)
            elif file_ext in ['.txt', '.md']:
                result = await self._preprocess_text_document(file_path, filename)
            else:
                # 알 수 없는 형식 - 텍스트로 처리 시도
                result = await self._preprocess_unknown_document(file_path, filename)

            processing_time = (datetime.now() - start_time).total_seconds()
            result.processing_time = processing_time

            logger.info(f"✅ [PREPROCESS] Completed for {filename}: {len(result.extracted_text)} chars extracted")

            return result

        except Exception as e:
            processing_time = (datetime.now() - start_time).total_seconds()
            logger.error(f"❌ [PREPROCESS] Failed for {filename}: {e}")

            return PreprocessingResult(
                success=False,
                extracted_text="",
                method_used="failed",
                metadata={"error": str(e)},
                processing_time=processing_time,
                char_count=0,
                error_message=str(e)
            )

    async def _preprocess_word_document(self, file_path: str, filename: str) -> PreprocessingResult:
        """Word 문서 전처리 - 다양한 방법 시도"""
        fallback_attempts = []

        # 방법 1: 향상된 python-docx 사용
        try:
            fallback_attempts.append("enhanced_python_docx")
            text = await self._extract_word_enhanced(file_path)
            if text.strip():
                return PreprocessingResult(
                    success=True,
                    extracted_text=text,
                    method_used="enhanced_python_docx",
                    metadata={"format": "docx", "method": "enhanced_python_docx"},
                    processing_time=0,
                    char_count=len(text),
                    fallback_attempts=fallback_attempts
                )
        except Exception as e:
            logger.warning(f"⚠️ [PREPROCESS] Enhanced python-docx failed: {e}")

        # 방법 2: 원시 XML 파싱
        try:
            fallback_attempts.append("raw_xml_parsing")
            text = await self._extract_word_raw_xml(file_path)
            if text.strip():
                return PreprocessingResult(
                    success=True,
                    extracted_text=text,
                    method_used="raw_xml_parsing",
                    metadata={"format": "docx", "method": "raw_xml_parsing"},
                    processing_time=0,
                    char_count=len(text),
                    fallback_attempts=fallback_attempts
                )
        except Exception as e:
            logger.warning(f"⚠️ [PREPROCESS] Raw XML parsing failed: {e}")

        # 방법 3: 기본 python-docx (단순)
        try:
            fallback_attempts.append("basic_python_docx")
            text = await self._extract_word_basic(file_path)
            if text.strip():
                return PreprocessingResult(
                    success=True,
                    extracted_text=text,
                    method_used="basic_python_docx",
                    metadata={"format": "docx", "method": "basic_python_docx"},
                    processing_time=0,
                    char_count=len(text),
                    fallback_attempts=fallback_attempts
                )
        except Exception as e:
            logger.warning(f"⚠️ [PREPROCESS] Basic python-docx failed: {e}")

        # 방법 4: 바이너리 텍스트 추출 시도
        try:
            fallback_attempts.append("binary_text_extraction")
            text = await self._extract_binary_text(file_path)
            if text.strip():
                return PreprocessingResult(
                    success=True,
                    extracted_text=text,
                    method_used="binary_text_extraction",
                    metadata={"format": "docx", "method": "binary_text_extraction", "warning": "텍스트 품질이 낮을 수 있음"},
                    processing_time=0,
                    char_count=len(text),
                    fallback_attempts=fallback_attempts
                )
        except Exception as e:
            logger.warning(f"⚠️ [PREPROCESS] Binary text extraction failed: {e}")

        # 모든 방법 실패
        return PreprocessingResult(
            success=False,
            extracted_text="",
            method_used="all_methods_failed",
            metadata={"format": "docx", "error": "모든 추출 방법 실패"},
            processing_time=0,
            char_count=0,
            error_message="DOCX 파일에서 텍스트를 추출할 수 없습니다.",
            fallback_attempts=fallback_attempts
        )

    async def _extract_word_enhanced(self, file_path: str) -> str:
        """향상된 Word 텍스트 추출"""
        try:
            from docx import Document
            from docx.oxml.table import CT_Tbl
            from docx.oxml.text.paragraph import CT_P
            from docx.table import _Cell, Table
            from docx.text.paragraph import Paragraph

            doc = Document(file_path)
            text_parts = []

            # 문서 메타데이터
            core_props = doc.core_properties
            if core_props.title:
                text_parts.append(f"제목: {core_props.title}")
            if core_props.subject:
                text_parts.append(f"주제: {core_props.subject}")
            if core_props.author:
                text_parts.append(f"작성자: {core_props.author}")

            if text_parts:
                text_parts.append("=" * 50)

            # 순서대로 요소 처리
            for element in doc.element.body:
                if element.tag.endswith('}p'):  # 단락
                    paragraph = Paragraph(element, doc)
                    if paragraph.text.strip():
                        # 스타일 정보 추가
                        style_name = paragraph.style.name
                        if style_name != 'Normal':
                            text_parts.append(f"[{style_name}] {paragraph.text}")
                        else:
                            text_parts.append(paragraph.text)

                elif element.tag.endswith('}tbl'):  # 표
                    table = Table(element, doc)
                    text_parts.append("\n--- 표 시작 ---")

                    for i, row in enumerate(table.rows):
                        row_data = []
                        for cell in row.cells:
                            cell_text = []
                            for para in cell.paragraphs:
                                if para.text.strip():
                                    cell_text.append(para.text.strip())
                            row_data.append(" ".join(cell_text) if cell_text else "")

                        if i == 0:  # 헤더
                            text_parts.append(f"헤더: {' | '.join(row_data)}")
                        else:
                            text_parts.append(f"행{i}: {' | '.join(row_data)}")

                    text_parts.append("--- 표 끝 ---\n")

            # 헤더/푸터 처리
            for section in doc.sections:
                if section.header:
                    header_text = self._extract_header_footer_text(section.header)
                    if header_text:
                        text_parts.insert(-1 if text_parts else 0, f"[헤더] {header_text}")

                if section.footer:
                    footer_text = self._extract_header_footer_text(section.footer)
                    if footer_text:
                        text_parts.append(f"[푸터] {footer_text}")

            return "\n".join(text_parts)

        except Exception as e:
            raise Exception(f"Enhanced Word extraction failed: {e}")

    def _extract_header_footer_text(self, header_or_footer) -> str:
        """헤더/푸터 텍스트 추출"""
        text_parts = []
        for paragraph in header_or_footer.paragraphs:
            if paragraph.text.strip():
                text_parts.append(paragraph.text.strip())
        return " ".join(text_parts)

    async def _extract_word_raw_xml(self, file_path: str) -> str:
        """원시 XML 파싱으로 Word 텍스트 추출"""
        try:
            import zipfile
            import xml.etree.ElementTree as ET

            text_parts = []

            with zipfile.ZipFile(file_path, 'r') as docx_zip:
                # document.xml 읽기
                if 'word/document.xml' in docx_zip.namelist():
                    with docx_zip.open('word/document.xml') as doc_xml:
                        content = doc_xml.read().decode('utf-8')

                        # XML 파싱
                        root = ET.fromstring(content)

                        # 네임스페이스 정의
                        namespaces = {
                            'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
                        }

                        # 모든 텍스트 요소 찾기
                        for text_elem in root.iter():
                            if text_elem.tag.endswith('}t'):
                                if text_elem.text:
                                    text_parts.append(text_elem.text)
                            elif text_elem.tag.endswith('}tab'):
                                text_parts.append('\t')
                            elif text_elem.tag.endswith('}br'):
                                text_parts.append('\n')

                # 추가로 headers와 footers 확인
                for file_name in docx_zip.namelist():
                    if 'word/header' in file_name or 'word/footer' in file_name:
                        try:
                            with docx_zip.open(file_name) as hf_xml:
                                hf_content = hf_xml.read().decode('utf-8')
                                hf_root = ET.fromstring(hf_content)

                                for text_elem in hf_root.iter():
                                    if text_elem.tag.endswith('}t') and text_elem.text:
                                        text_parts.append(f"[헤더/푸터] {text_elem.text}")
                        except:
                            continue

            return " ".join(text_parts) if text_parts else ""

        except Exception as e:
            raise Exception(f"Raw XML extraction failed: {e}")

    async def _extract_word_basic(self, file_path: str) -> str:
        """기본 Word 텍스트 추출"""
        try:
            from docx import Document

            doc = Document(file_path)
            text_parts = []

            # 단락 텍스트
            for para in doc.paragraphs:
                if para.text.strip():
                    text_parts.append(para.text)

            # 표 텍스트
            for table in doc.tables:
                for row in table.rows:
                    row_text = []
                    for cell in row.cells:
                        if cell.text.strip():
                            row_text.append(cell.text.strip())
                    if row_text:
                        text_parts.append(" | ".join(row_text))

            return "\n".join(text_parts)

        except Exception as e:
            raise Exception(f"Basic Word extraction failed: {e}")

    async def _extract_binary_text(self, file_path: str) -> str:
        """바이너리 파일에서 텍스트 추출 시도"""
        try:
            with open(file_path, 'rb') as f:
                content = f.read()

            # 다양한 인코딩으로 디코딩 시도
            encodings = ['utf-8', 'cp949', 'euc-kr', 'latin-1', 'cp1252']

            for encoding in encodings:
                try:
                    decoded = content.decode(encoding, errors='ignore')
                    # 텍스트로 보이는 부분만 추출 (알파벳, 숫자, 한글, 공백, 구두점)
                    import re
                    text_parts = re.findall(r'[a-zA-Z0-9가-힣\s.,!?;:\'\"()-]+', decoded)
                    clean_text = ' '.join(part.strip() for part in text_parts if len(part.strip()) > 2)

                    if len(clean_text) > 50:  # 충분한 텍스트가 추출된 경우
                        return clean_text

                except UnicodeDecodeError:
                    continue

            return ""

        except Exception as e:
            raise Exception(f"Binary text extraction failed: {e}")

    async def _preprocess_pdf_document(self, file_path: str, filename: str) -> PreprocessingResult:
        """PDF 문서 전처리"""
        # 기존 PDF 처리 로직 유지하면서 향상
        try:
            import PyPDF2

            text_parts = []
            with open(file_path, 'rb') as file:
                pdf_reader = PyPDF2.PdfReader(file)

                for page_num, page in enumerate(pdf_reader.pages):
                    page_text = page.extract_text()
                    if page_text.strip():
                        text_parts.append(f"--- 페이지 {page_num + 1} ---\n{page_text}")

            extracted_text = "\n".join(text_parts)

            return PreprocessingResult(
                success=True,
                extracted_text=extracted_text,
                method_used="PyPDF2",
                metadata={"format": "pdf", "pages": len(pdf_reader.pages)},
                processing_time=0,
                char_count=len(extracted_text),
                fallback_attempts=["PyPDF2"]
            )

        except Exception as e:
            return PreprocessingResult(
                success=False,
                extracted_text="",
                method_used="failed",
                metadata={"format": "pdf", "error": str(e)},
                processing_time=0,
                char_count=0,
                error_message=f"PDF 처리 실패: {str(e)}"
            )

    async def _preprocess_powerpoint_document(self, file_path: str, filename: str) -> PreprocessingResult:
        """PowerPoint 문서 전처리"""
        try:
            from pptx import Presentation

            prs = Presentation(file_path)
            text_parts = []

            for slide_num, slide in enumerate(prs.slides):
                text_parts.append(f"\n--- 슬라이드 {slide_num + 1} ---")

                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        text_parts.append(shape.text)

            extracted_text = "\n".join(text_parts)

            return PreprocessingResult(
                success=True,
                extracted_text=extracted_text,
                method_used="python_pptx",
                metadata={"format": "pptx", "slides": len(prs.slides)},
                processing_time=0,
                char_count=len(extracted_text),
                fallback_attempts=["python_pptx"]
            )

        except Exception as e:
            return PreprocessingResult(
                success=False,
                extracted_text="",
                method_used="failed",
                metadata={"format": "pptx", "error": str(e)},
                processing_time=0,
                char_count=0,
                error_message=f"PowerPoint 처리 실패: {str(e)}"
            )

    async def _preprocess_excel_document(self, file_path: str, filename: str) -> PreprocessingResult:
        """Excel 문서 전처리"""
        try:
            import openpyxl

            workbook = openpyxl.load_workbook(file_path, data_only=True)
            text_parts = []

            for sheet_name in workbook.sheetnames:
                sheet = workbook[sheet_name]
                text_parts.append(f"\n--- 시트: {sheet_name} ---")

                for row in sheet.iter_rows(values_only=True):
                    row_text = "\t".join(str(cell) if cell is not None else "" for cell in row)
                    if row_text.strip():
                        text_parts.append(row_text)

            extracted_text = "\n".join(text_parts)

            return PreprocessingResult(
                success=True,
                extracted_text=extracted_text,
                method_used="openpyxl",
                metadata={"format": "xlsx", "sheets": len(workbook.sheetnames)},
                processing_time=0,
                char_count=len(extracted_text),
                fallback_attempts=["openpyxl"]
            )

        except Exception as e:
            return PreprocessingResult(
                success=False,
                extracted_text="",
                method_used="failed",
                metadata={"format": "xlsx", "error": str(e)},
                processing_time=0,
                char_count=0,
                error_message=f"Excel 처리 실패: {str(e)}"
            )

    async def _preprocess_text_document(self, file_path: str, filename: str) -> PreprocessingResult:
        """텍스트 문서 전처리"""
        encodings = ['utf-8', 'cp949', 'euc-kr', 'latin-1']

        for encoding in encodings:
            try:
                with open(file_path, 'r', encoding=encoding) as file:
                    text = file.read()

                return PreprocessingResult(
                    success=True,
                    extracted_text=text,
                    method_used=f"text_file_{encoding}",
                    metadata={"format": "text", "encoding": encoding},
                    processing_time=0,
                    char_count=len(text),
                    fallback_attempts=[f"text_file_{encoding}"]
                )

            except UnicodeDecodeError:
                continue

        # 모든 인코딩 실패시 에러 무시하고 읽기
        try:
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as file:
                text = file.read()

            return PreprocessingResult(
                success=True,
                extracted_text=text,
                method_used="text_file_utf8_ignore",
                metadata={"format": "text", "encoding": "utf-8", "warning": "일부 문자가 손실될 수 있음"},
                processing_time=0,
                char_count=len(text),
                fallback_attempts=["text_file_utf8_ignore"]
            )

        except Exception as e:
            return PreprocessingResult(
                success=False,
                extracted_text="",
                method_used="failed",
                metadata={"format": "text", "error": str(e)},
                processing_time=0,
                char_count=0,
                error_message=f"텍스트 파일 처리 실패: {str(e)}"
            )

    async def _preprocess_unknown_document(self, file_path: str, filename: str) -> PreprocessingResult:
        """알 수 없는 형식의 문서 전처리"""
        logger.info(f"🔍 [PREPROCESS] Unknown format, attempting text extraction for {filename}")

        # 텍스트로 처리 시도
        return await self._preprocess_text_document(file_path, filename)

# 싱글톤 인스턴스
_preprocessor_instance = None

def get_advanced_preprocessor() -> AdvancedDocumentPreprocessor:
    """고급 전처리기 인스턴스 가져오기"""
    global _preprocessor_instance
    if _preprocessor_instance is None:
        _preprocessor_instance = AdvancedDocumentPreprocessor()
    return _preprocessor_instance